"""
AI-powered column detection using Google Gemini API (free tier available).
Sends a sample of the sheet (headers + first rows) to Gemini Flash,
asks it to identify which column index maps to which field,
then uses that mapping to extract records.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

from parser import (
    _to_date, _str, _find_header_end,
    ESP_SHEETS, parse_esp,
    _is_submission_control,
)

# ── fields we want Gemini to identify ────────────────────────────────────────
FIELD_DESCRIPTIONS = {
    "item_name":  "材料/設備名稱 (material or equipment name)",
    "po_no":      "PO編號/文件編碼 (purchase order number or document code, alphanumeric)",
    "tag_no":     "設備Tag號/設備編號 (equipment tag number)",
    "sub_vendor": "次廠商/供應商名稱 (sub-vendor or supplier name)",
    "ros":        "需求到料日期/預定進場日期 (required on site date — must be a date column)",
    "eta":        "預計到料/預定廠驗日期 (expected delivery or factory inspection date — must be a date column)",
    "ata":        "實際到料日期 (actual delivery date — must be a date column)",
    "qty":        "數量 (quantity)",
}

SYSTEM_PROMPT = """\
You are a data extraction assistant. Analyse a spreadsheet table (given as JSON rows)
and identify which 0-based column index contains each requested field.
Return ONLY a valid JSON object — no markdown, no explanation.

Rules:
- For date fields (ros, eta, ata): column must contain calendar dates, NOT document codes.
- po_no must contain alphanumeric codes (e.g. ED2024-001), NOT dates.
- If a field does not exist, use null.
- Output format: {"item_name":<int|null>,"po_no":<int|null>,"tag_no":<int|null>,
  "sub_vendor":<int|null>,"ros":<int|null>,"eta":<int|null>,"ata":<int|null>,"qty":<int|null>}
"""


def _sheet_to_sample(rows: list[tuple], max_rows: int = 10) -> str:
    sample = []
    for row in rows[:max_rows]:
        sample.append([_str(c).split("\n")[0] for c in row])
    return json.dumps(sample, ensure_ascii=False)


GEMINI_URL = (
    "https://generativelanguage.googleapis.com/v1beta/models/"
    "gemini-2.0-flash:generateContent?key={api_key}"
)

def _call_gemini(api_key: str, table_sample: str, sheet_name: str) -> dict | None:
    import urllib.request

    prompt = (
        f"{SYSTEM_PROMPT}\n\n"
        f"Sheet: {sheet_name}\n\n"
        f"Table sample (JSON rows, 0-based columns):\n{table_sample}\n\n"
        "Fields to find:\n"
        + "\n".join(f"- {k}: {v}" for k, v in FIELD_DESCRIPTIONS.items())
        + "\n\nReturn only the JSON object."
    )

    payload = json.dumps({
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0, "maxOutputTokens": 256},
    }).encode()

    req = urllib.request.Request(
        GEMINI_URL.format(api_key=api_key),
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            body = json.loads(resp.read())
        raw = body["candidates"][0]["content"]["parts"][0]["text"].strip()
    except Exception:
        return None

    raw = re.sub(r"^```[a-z]*\n?", "", raw)
    raw = re.sub(r"\n?```$", "", raw)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return None


def _extract_records(rows, col_map, sheet_name, file_name) -> list[dict]:
    item_col   = col_map.get("item_name")
    if item_col is None:
        return []
    tag_col    = col_map.get("tag_no")
    po_col     = col_map.get("po_no")
    vendor_col = col_map.get("sub_vendor")
    ros_col    = col_map.get("ros")
    eta_col    = col_map.get("eta")
    ata_col    = col_map.get("ata")
    qty_col    = col_map.get("qty")

    data_start = _find_header_end(rows, {"item": item_col})
    records = []
    seen: set = set()
    current_tag = ""

    for row in rows[data_start:]:
        def g(ci):
            return row[ci] if ci is not None and len(row) > ci else None

        raw_tag = _str(g(tag_col))
        if raw_tag:
            current_tag = raw_tag.split("\n")[0]

        item_name = _str(g(item_col)).split("\n")[0].strip()
        if not item_name:
            continue
        if re.search(
            r"合計|小計|total|subtotal|^(文件|設備名稱|item|description|材料名稱|編碼)$",
            item_name, re.I,
        ):
            continue

        key = (sheet_name, item_name)
        if key in seen:
            continue
        seen.add(key)

        eta = _to_date(g(eta_col))
        ata = _to_date(g(ata_col))
        ros = _to_date(g(ros_col))

        records.append({
            "source_file":       file_name,
            "sheet":             sheet_name,
            "mr_no":             "",
            "mr_name":           "",
            "po_no":             _str(g(po_col)).split("\n")[0],
            "tag_no":            current_tag,
            "vendor":            "",
            "delivery_location": "",
            "ros":               ros,
            "item_name":         item_name,
            "qty":               _str(g(qty_col)),
            "sub_vendor":        _str(g(vendor_col)).split("\n")[0],
            "eta":               eta,
            "ata":               ata,
            "effective_delivery": ata or eta or ros,
            "sub_order_received": None,
            "record_type":       "AI解析",
        })
    return records


# ── Public entry point ────────────────────────────────────────────────────────

def parse_with_ai(path: Path, api_key: str) -> tuple[list[dict], list[dict]]:
    """
    Parse any Excel with Gemini for column detection.
    Returns (records, mapping_log).
    """
    ext = path.suffix.lower()
    if ext not in (".xlsx", ".xls"):
        return [], []

    wb = openpyxl.load_workbook(str(path), data_only=True)

    # ESP format is already handled perfectly — skip AI for it
    if set(wb.sheetnames) & ESP_SHEETS:
        wb.close()
        return parse_esp(path), []

    all_records = []
    mapping_log = []

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows or ws.max_row < 3:
            continue

        non_empty = [r for r in rows[:12] if any(c for c in r)]
        sample    = _sheet_to_sample(non_empty, max_rows=10)

        raw_map = _call_gemini(api_key, sample, ws.title)
        if not raw_map:
            mapping_log.append({"sheet": ws.title, "status": "AI無法解析", "mapping": {}})
            continue

        col_map = {k: v for k, v in raw_map.items() if v is not None}
        ok = col_map.get("item_name") is not None
        mapping_log.append({
            "sheet":   ws.title,
            "status":  "成功" if ok else "找不到材料名稱欄",
            "mapping": col_map,
        })
        if not ok:
            continue

        all_records.extend(_extract_records(rows, col_map, ws.title, path.name))

    wb.close()
    return all_records, mapping_log
